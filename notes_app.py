import os
import sys
import tkinter as tk
from datetime import datetime
from tkinter import messagebox, ttk

from openpyxl import Workbook, load_workbook

# ملف الإكسل يُحفظ بجانب البرنامج
if getattr(sys, "frozen", False):
    BASE = os.path.dirname(sys.executable)  # عند التشغيل كملف exe
else:
    BASE = os.path.dirname(os.path.abspath(__file__))
FILE = os.path.join(BASE, "ملاحظاتي.xlsx")


def load_book():
    if os.path.exists(FILE):
        return load_workbook(FILE)
    wb = Workbook()
    ws = wb.active
    ws.title = "الملاحظات"
    ws.sheet_view.rightToLeft = True
    ws.append(["التاريخ", "الوقت", "الملاحظة"])
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 80
    return wb


def save_note():
    text = box.get("1.0", tk.END).strip()
    if not text:
        messagebox.showwarning("تنبيه", "اكتب ملاحظة أولاً")
        return
    now = datetime.now()
    try:
        wb = load_book()
        ws = wb.active
        ws.append([now.strftime("%Y-%m-%d"), now.strftime("%H:%M:%S"), text])
        wb.save(FILE)
    except PermissionError:
        messagebox.showerror("خطأ", "أغلق ملف الإكسل أولاً ثم أعد المحاولة")
        return
    box.delete("1.0", tk.END)
    refresh_list()
    status.config(text="تم الحفظ ✓  " + now.strftime("%H:%M:%S"))


def refresh_list():
    listbox.delete(0, tk.END)
    if not os.path.exists(FILE):
        return
    ws = load_workbook(FILE).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for d, t, n in reversed(rows[-20:]):
        listbox.insert(tk.END, f"{d}  {t}  —  {str(n).replace(chr(10), ' ')[:60]}")


def show_all():
    if not os.path.exists(FILE):
        messagebox.showinfo("ملاحظاتي", "لا توجد ملاحظات محفوظة بعد")
        return
    try:
        ws = load_workbook(FILE).active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
    except PermissionError:
        messagebox.showerror("خطأ", "أغلق ملف الإكسل أولاً ثم أعد المحاولة")
        return

    win = tk.Toplevel(root)
    win.title(f"كل الملاحظات ({len(rows)})")
    win.geometry("720x460")

    frame = tk.Frame(win)
    frame.pack(fill="both", expand=True, padx=10, pady=10)

    tree = ttk.Treeview(frame, columns=("date", "time", "note"), show="headings")
    tree.heading("date", text="التاريخ")
    tree.heading("time", text="الوقت")
    tree.heading("note", text="الملاحظة")
    tree.column("date", width=100, anchor="center", stretch=False)
    tree.column("time", width=80, anchor="center", stretch=False)
    tree.column("note", width=500, anchor="e")

    scroll = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=scroll.set)
    scroll.pack(side="left", fill="y")
    tree.pack(side="right", fill="both", expand=True)

    # الأحدث أولاً
    for d, t, n in reversed(rows):
        tree.insert("", tk.END, values=(d, t, str(n).replace("\n", " ")))

    def show_full(_event=None):
        sel = tree.selection()
        if not sel:
            return
        d, t, n = tree.item(sel[0], "values")
        messagebox.showinfo(f"{d}  {t}", n)

    tree.bind("<Double-1>", show_full)


root = tk.Tk()
root.title("ملاحظاتي اليومية")
root.geometry("560x560")

tk.Label(root, text="اكتب ملاحظتك:", font=("Arial", 13)).pack(anchor="e", padx=12, pady=(12, 4))
box = tk.Text(root, height=8, font=("Arial", 13), wrap="word")
box.pack(fill="x", padx=12)
box.tag_configure("rtl", justify="right")

buttons = tk.Frame(root)
buttons.pack(pady=10)
tk.Button(buttons, text="حفظ الملاحظة", font=("Arial", 13, "bold"), bg="#2e7d32", fg="white",
          command=save_note).pack(side="right", padx=6)
tk.Button(buttons, text="عرض كل الملاحظات", font=("Arial", 13), bg="#1565c0", fg="white",
          command=show_all).pack(side="right", padx=6)
status = tk.Label(root, text="", fg="#2e7d32")
status.pack()

tk.Label(root, text="آخر الملاحظات:", font=("Arial", 12)).pack(anchor="e", padx=12, pady=(10, 2))
listbox = tk.Listbox(root, font=("Arial", 11))
listbox.pack(fill="both", expand=True, padx=12, pady=(0, 12))

refresh_list()
root.mainloop()
